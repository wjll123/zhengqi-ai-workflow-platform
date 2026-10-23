"""平台工作流的 Excel 导出和历史结果导入。"""
from pathlib import Path
import math
import unicodedata
from tempfile import NamedTemporaryFile
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .errors import QueryError
from .sources import _path, normalize_name
from .fields import SCHEMA_VERSION, FIELD_SCOPE, SHARE_FIELDS, PEOPLE_FIELDS, DETAIL_HEADERS, is_full

LABELS = {"success": "成功", "failed": "失败", "pending": "待查询", "blocked": "待验证", "querying": "查询中"}


def safe(value):
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value


def read_previous(filename, names, root):
    path = _path(filename, root)
    if path.suffix.lower() != ".xlsx":
        raise QueryError("历史结果需为本功能生成的 Excel。")
    try:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if not {"股东信息", "主要人物", "查询状态"}.issubset(book.sheetnames):
                raise QueryError("历史结果必须包含股东信息、主要人物、查询状态三张表。")
            rows = list(book["查询状态"].values)
            if not rows or list(rows[0][:6]) != ["公司名称", "状态", "股东行数", "主要人物行数", "说明", "来源链接"]:
                raise QueryError("历史结果查询状态表格式不正确。")
            if [normalize_name(row[0]) for row in rows[1:]] != [normalize_name(name) for name in names]:
                raise QueryError("历史结果与当前名单不一致，请选择同一份名单。")
            successes = {}
            status_fields = {name: index for index, name in enumerate(rows[0])}
            for name, row in zip(names, rows[1:]):
                if row[1] not in LABELS.values():
                    raise QueryError("历史结果存在未知状态。")
                if row[2] is not None and row[3] is not None:
                    parsed = urlsplit(str(row[5] or ""))
                    if parsed.scheme != "https" or parsed.hostname not in {"www.qcc.com", "m.qcc.com", "qcc.com"} or not parsed.path.startswith("/firm/"):
                        raise QueryError("历史结果的来源链接无效。")
                    version = row[status_fields["字段版本"]] if "字段版本" in status_fields else None
                    successes[normalize_name(name)] = {"name": name, "status": "success" if row[1] == "成功" else "pending", "error": "",
                        "queried_at": str(row[6] or "") if len(row) > 6 else "",
                        "counts": {"shareholders": row[2], "people": row[3],
                                   "contributions": row[status_fields["出资明细行数"]] if "出资明细行数" in status_fields else 0},
                        "note": str(row[4] or ""),
                        "data": {"company": str(row[8] or name) if len(row) > 8 else name, "url": row[5], "credit_code": str(row[7] or "") if len(row) > 7 else "",
                                 "schema_version": version,
                                 "field_scope": FIELD_SCOPE if version == SCHEMA_VERSION else "旧版基础字段",
                                 "additional_details": row[status_fields["补充详情提示"]] if "补充详情提示" in status_fields else ""}}
            for title, kind in (("股东信息", "shareholders"), ("主要人物", "people"), ("股东出资明细", "contributions")):
                if title not in book.sheetnames:
                    if any(is_full(item["data"]) for item in successes.values()):
                        raise QueryError("完整字段版本缺少股东出资明细表。")
                    continue
                detail_rows = book[title].values
                headers = next(detail_rows, ())
                if not headers or headers[0] != "公司名称" or headers[-1] != "来源链接":
                    raise QueryError(f"历史结果的{title}表头不正确。")
                for item in successes.values():
                    item["data"][kind] = {"headers": list(headers[1:-1]), "rows": []}
                for row in detail_rows:
                    key = normalize_name(row[0])
                    if key not in successes or row[-1] != successes[key]["data"]["url"]:
                        raise QueryError("历史结果的明细与状态或来源链接不一致。")
                    successes[key]["data"][kind]["rows"].append(list(row[1:-1]))
            for item in successes.values():
                for kind in ("shareholders", "people", "contributions"):
                    if kind not in item["data"]:
                        continue
                    if len(item["data"][kind]["rows"]) != item["counts"][kind]:
                        raise QueryError("历史结果的明细条数与状态表不一致。")
                if is_full(item["data"]):
                    for kind, fields in (("shareholders", SHARE_FIELDS), ("people", PEOPLE_FIELDS)):
                        if not {label for _, label in fields}.issubset(item["data"][kind]["headers"]):
                            raise QueryError("历史结果缺少完整字段列，请重新补齐。")
                item.pop("counts")
            return {item["name"]: item for item in successes.values()}
        finally:
            book.close()
    except QueryError:
        raise
    except Exception as exc:
        raise QueryError("无法读取历史查询结果，请检查文件。") from exc


def export_job(job, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    retained = [row for row in job["rows"] if row.get("data")]
    for title, kind in (("股东信息", "shareholders"), ("主要人物", "people"), ("股东出资明细", "contributions")):
        sheet = book.create_sheet(title)
        base_fields = [label for _, label in SHARE_FIELDS] if kind == "shareholders" else [label for _, label in PEOPLE_FIELDS] if kind == "people" else DETAIL_HEADERS
        extra_fields = [field for row in retained for field in row["data"].get(kind, {}).get("headers", [])
                        if field not in {"资料范围", "采集时间", "出资日期"}]
        fields = list(dict.fromkeys(base_fields + extra_fields))
        sheet.append(["公司名称", *fields, "资料范围", "采集时间", "来源链接"])
        for row in retained:
            data = row["data"]
            for details in data.get(kind, {}).get("rows", []):
                values = dict(zip(data[kind]["headers"], details))
                if "出资日期" in values and "认缴出资日期" not in values:
                    values["认缴出资日期"] = values["出资日期"]
                scope = FIELD_SCOPE if is_full(data) else "旧版基础字段，待补齐"
                sheet.append([safe(row["name"]), *(safe(values.get(field)) for field in fields), scope, row.get("queried_at", ""), data["url"]])
        if "持股比例" in fields:
            column = fields.index("持股比例") + 2
            for index in range(2, sheet.max_row + 1):
                sheet.cell(index, column).number_format = "0.0000%"
    sheet = book.create_sheet("查询状态")
    sheet.append(["公司名称", "状态", "股东行数", "主要人物行数", "说明", "来源链接", "查询时间", "统一社会信用代码", "企查查登记名称",
                  "字段范围", "字段版本", "记录校验", "补充详情提示", "出资明细行数"])
    for row in job["rows"]:
        data = row.get("data") or {}
        sheet.append([safe(row["name"]), LABELS[row["status"]],
                      len(data.get("shareholders", {}).get("rows", [])) if data else None,
                      len(data.get("people", {}).get("rows", [])) if data else None,
                      safe(row.get("error") or row.get("note", "")), data.get("url", ""),
                      row.get("queried_at", ""), data.get("credit_code", ""), safe(data.get("company", "")),
                      (FIELD_SCOPE if is_full(data) else "旧版基础字段，待补齐") if data else "尚未取得资料",
                      data.get("schema_version", 1) if data else None,
                      "股东、主要人员条数与页面总数一致" if is_full(data) else "待补齐字段后复核",
                      safe(data.get("additional_details", "")), len(data.get("contributions", {}).get("rows", [])) if data else None])
    for row in sheet.iter_rows(min_row=2):
        row[7].number_format = "@"
    notes = book.create_sheet("字段说明")
    for row in [
        ["项目", "说明"],
        ["资料范围", "保存企业详情页提供的全部股东、主要人员记录与业务字段，逐家核对页面总数；并非人物独立档案、持股穿透或历史变更的全部页面。"],
        ["补充详情提示", "页面显示人物简介或持股详情入口，但正文不在企业详情页时，在对应行明确标明未采集。"],
        ["未披露", "源页面字段为空、null 或空列表，不能解读为零；数值 0 会原样保留。"],
        ["旧版基础字段", "补齐期间仍保留旧结果，并单独标注资料范围，不能视作已补齐。"],
        ["出资明细", "认缴、实缴的每一期独立一行。页面未返回分期列表时不伪造明细；认缴、实缴汇总见股东信息。金额沿用页面万元口径。"],
        ["日期", "认缴、实缴日期沿用页面原文；入股、任职日期从页面时间戳按北京时间转换，未提供时标为未披露。"],
        ["证照号码", "企业营业执照号码按原文保留；自然人证件号码脱敏。"],
        ["标签补充信息", "保留页面标签数据；BPL、IsBP 等源字段不自行推断含义。"],
    ]:
        notes.append(row)
    for sheet in book:
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="173A52")
            cell.font = Font(name="等线", bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for index in range(1, sheet.max_column + 1):
            title = sheet.cell(1, index).value
            sheet.column_dimensions[get_column_letter(index)].width = 42 if title in ("公司名称", "股东名称", "来源链接", "企查查登记名称") else 60 if title in ("说明", "补充详情提示", "其他披露字段") else 24
        for row in sheet.iter_rows(min_row=2):
            lines = 1
            for cell in row:
                cell.font = Font(name="等线", size=11, color="1B3546")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F5F7")
                if "号码" in str(sheet.cell(1, cell.column).value):
                    cell.number_format = "@"
                width = sheet.column_dimensions[cell.column_letter].width
                for line in str(cell.value or "").splitlines():
                    units = sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line)
                    lines = max(lines, math.ceil(units / max(1, width - 2)))
                if sheet.cell(1, cell.column).value == "来源链接" and cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(name="等线", size=11, color="1967A3", underline="single")
            sheet.row_dimensions[row[0].row].height = min(409, max(36, lines * 16 + 8))
    notes.column_dimensions["B"].width = 100
    for index in range(2, notes.max_row + 1):
        notes.row_dimensions[index].height = 48
    destination = output_dir / f"企查查-批量查询-{job['id']}.xlsx"
    temp = None
    try:
        with NamedTemporaryFile(dir=output_dir, suffix=".xlsx", delete=False) as stream:
            temp = Path(stream.name)
        book.save(temp)
        temp.replace(destination)
    finally:
        book.close()
        if temp:
            temp.unlink(missing_ok=True)
    return destination
