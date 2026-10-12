"""表格读写层

处理平台遇到的所有"坑"格式：
- BIFF 格式 .xls（WPS 导出的 OLE2）→ 必须用 xlrd，openpyxl 读不了
- GBK 编码的 CSV → 必须指定编码
- 单元格尾部的 \\t 制表符 → 统一清洗
"""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook


def _clean(value: Any) -> Any:
    """清洗单元格值：去掉首尾空白与制表符"""
    if isinstance(value, str):
        return value.strip().strip("\t").strip()
    return value


def read_csv_gbk(path: str | Path) -> list[list[Any]]:
    """读取 GBK 编码的 CSV（智慧中台调度工单导出格式）"""
    path = Path(path)
    for enc in ("gbk", "gb18030", "utf-8-sig", "utf-8"):
        try:
            with open(path, encoding=enc, newline="") as f:
                rows = list(csv.reader(f))
            return [[_clean(c) for c in row] for row in rows]
        except (UnicodeDecodeError, LookupError):
            continue
    raise ValueError(f"无法识别 CSV 编码: {path}")


def read_xls_biff(path: str | Path, sheet_index: int = 0) -> list[list[Any]]:
    """读取 BIFF 格式 .xls（xlrd）"""
    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_index(sheet_index)
    out: list[list[Any]] = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            v = sheet.cell_value(r, c)
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            row.append(_clean(v))
        out.append(row)
    return out


def read_xlsx(path: str | Path, sheet: str | None = None) -> list[list[Any]]:
    """读取 .xlsx"""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def read_any(path: str | Path, sheet: str | None = None) -> list[list[Any]]:
    """自动识别格式读取（.xls / .csv / .xlsx）"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_gbk(path)
    if suffix == ".xls":
        return read_xls_biff(path)
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx(path, sheet)
    raise ValueError(f"不支持的格式: {suffix}")


def upsert_sheet(
    path: str | Path,
    sheet_name: str,
    rows: list[list[Any]],
    sheet_order: list[str] | None = None,
) -> Path:
    """向已有 .xlsx 追加/替换一个 sheet，保留其他 sheet 不动。

    用于「多个任务合成一个输出文件的多张子表」的场景 ——
    例如 1-接单情况模板.xlsx 由 4 个任务各写一个 sheet。

    sheet_order: 期望的 sheet 顺序（存在的才排，不存在忽略），其余按原顺序追加在后面。
    """
    from openpyxl import Workbook, load_workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        wb = load_workbook(path)
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
    else:
        wb = Workbook()
        wb.remove(wb.active)

    ws = wb.create_sheet(sheet_name)
    for row in rows:
        ws.append(list(row))

    if sheet_order:
        want = [s for s in sheet_order if s in wb.sheetnames]
        rest = [s for s in wb.sheetnames if s not in want]
        wb._sheets = [wb[s] for s in (want + rest)]

    wb.save(path)
    wb.close()
    return path


def write_xlsx(
    path: str | Path,
    sheets: dict[str, list[list[Any]]],
    number_formats: dict[tuple[str, int], str] | None = None,
) -> Path:
    """写出 .xlsx（支持按列设置数字格式）

    number_formats: {(sheet_name, col_index_0based): '0.00'} 形式
    """
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(list(row))
        if number_formats:
            for (sname, col), fmt in number_formats.items():
                if sname != name:
                    continue
                letter = ws.cell(row=1, column=col + 1).column_letter
                for r in range(1, len(rows) + 1):
                    ws[f"{letter}{r}"].number_format = fmt

    wb.save(path)
    wb.close()
    return path


def sheet_names(path: str | Path) -> list[str]:
    """列出所有 sheet 名"""
    path = Path(path)
    if path.suffix.lower() == ".xls":
        return xlrd.open_workbook(str(path)).sheet_names()
    wb = load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def cell_to_text(value: Any) -> str:
    """把单元格值转为文本（用于比对）"""
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(round(value, 10))
    return str(value).strip()
