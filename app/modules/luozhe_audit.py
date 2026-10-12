"""客户落责月度工作流的四文件只读核对。"""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from ..core.config import INBOX_DIR, OUTPUT_DIR
from ..core.runner import TaskResult, registry
from .data_luozhe import (CITY_ORDER, _norm_grow, _parse_stat_sheet, _pick_province,
                          _ppt_full_cells, _ppt_simple_cells)


SOURCE_SHEETS = (
    ("地市总体落责", "全量统计", 5, 33, 120),
    ("集团规上", "集团规上统计", 5, 35, 48),
    ("专精特新", "专精特新统计", 5, 35, 48),
)
MATERIALS = (
    ("全量", "全量收入正增长统计", 48, 15, 5, ("卓越", "战略", "核心"), {2, 3, 5, 7, 8, 10, 12, 13, 15}),
    ("集团规上", "集团规上统计", 27, 12, 4, ("战略", "核心", "小计"), {2, 4, 6, 8, 10, 12}),
    ("专精特新", "专精特新统计", 27, 12, 4, ("战略", "核心", "小计"), {2, 4, 6, 8, 10, 12}),
)
SOURCE_FIELDS = ((2, 3), (3, 5)) + tuple((c, c + 2) for c in range(4, 20))


def _input(name: str, suffix: str, inbox: Path) -> Path:
    if not name or Path(name).name != name or Path(name).suffix.lower() != suffix:
        raise ValueError(f"请选择有效的 {suffix} 文件")
    path = inbox / name
    if not path.is_file():
        raise ValueError(f"文件不存在：{name}")
    return path


def _equal(a, b) -> bool:
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-9
    return str(a).strip() == str(b).strip()


def _number(value) -> Decimal:
    return Decimal(str(value).strip().replace("%", ""))


def _compare_source(source, process, source_sheet, process_sheet, first_source, first_process, count):
    left, right = source[source_sheet], process[process_sheet]
    differences = []
    checked = 0
    for offset in range(count):
        source_row, process_row = first_source + offset, first_process + offset
        for sc, pc in ((1, 2),) + SOURCE_FIELDS:
            a, b = left.cell(source_row, sc), right.cell(process_row, pc)
            if sc != 1 and a.value is None and b.value is None:
                continue
            checked += 1
            if not _equal(a.value, b.value):
                differences.append((source_sheet, a.coordinate, process_sheet, b.coordinate, a.value, b.value))
    return checked, differences


def _compare_growth(source, process):
    left, right = source.worksheets[0], process["全量收入正增长统计"]
    differences = []
    checked = 0
    for offset in range(23):
        for sc in range(1, 8):
            a, b = left.cell(2 + offset, sc), right.cell(34 + offset, sc + 35)
            checked += 1
            if not _equal(a.value, b.value):
                differences.append((a.coordinate, b.coordinate, a.value, b.value))
    return checked, differences


def _compare_material(process, deck, config, page):
    label, sheet, start, ncols, block, grades, rates = config
    ws = process[sheet]
    table_shape = next((shape for shape in deck.slides[page].shapes if shape.has_table), None)
    if table_shape is None:
        raise ValueError(f"月报第 {page + 1} 页缺少可编辑表格")
    table = table_shape.table
    if len(table.rows) != 25 or len(table.columns) != ncols + 1:
        raise ValueError(f"月报第 {page + 1} 页表格结构与流程模板不符")
    differences = []
    names = []
    for offset in range(23):
        row = 7 + offset
        process_name = str(ws.cell(row, start).value or "").strip()
        slide_name = table.cell(offset + 2, 0).text.strip()
        aliases = {"省政企": "政企", "省直管客户": "省管客户"}
        if aliases.get(process_name, process_name) != slide_name:
            names.append((process_name, slide_name, row, offset + 3))
        for j in range(1, ncols + 1):
            cell = ws.cell(row, start + j)
            raw = table.cell(offset + 2, j).text.strip()
            try:
                expected = _number(cell.value)
                actual = _number(raw)
                # 月报比率只显示一位小数；允许 0.05 个百分点的显示误差。
                different = abs(expected - actual) > (Decimal("0.05") if j in rates else Decimal("0"))
            except (InvalidOperation, TypeError):
                different = str(cell.value or "").strip() != raw
            if different:
                differences.append((process_name, grades[(j - 1) // block],
                                    table.cell(1, j).text.strip(), cell.coordinate,
                                    cell.value, raw, offset + 3))
    return differences, names


def _compare_material_source(source, growth, process, config, page):
    """独立用来源表重算 PPT 材料，防止缓存公式或搬运区之外出错。"""
    _, sheet, start, ncols, _, _, _ = config
    source_sheet = SOURCE_SHEETS[page][0]
    source_rows = [list(row) for row in source[source_sheet].iter_rows(max_col=19, values_only=True)]
    stats = _parse_stat_sheet(source_rows)
    growth_rows = [list(row) for row in growth.worksheets[0].iter_rows(max_col=7, values_only=True)]
    growth_data = _norm_grow(growth_rows)
    records = [("全省", _pick_province(stats))] + [(city, stats.get(city, {})) for city in CITY_ORDER]
    differences = []
    for offset, (city, stat) in enumerate(records):
        if not stat:
            differences.append((city, "整行缺失", "", ""))
            continue
        expected = (_ppt_full_cells(stat, growth_data.get(city, {})) if page == 0
                    else _ppt_simple_cells(stat))
        for j in range(ncols):
            cell = process[sheet].cell(7 + offset, start + 1 + j)
            a, b = expected[j], cell.value
            if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                mismatch = abs(a - b) > 1e-7
            else:
                mismatch = not _equal(a, b)
            if mismatch:
                differences.append((city, cell.coordinate, a, b))
    return differences


@registry.register(
    key="monthly_luozhe_audit",
    name="客户落责月度工作流对比",
    module="数据分析",
    description="按 7 步流程核对两份来源表、处理工作簿和三页月报，输出差异明细",
    source_type="auto",
    source_note="上传并选择两份月度来源表、对应的落责材料数据处理工作簿和三页客户落责月报。四份文件应属于同一统计期。平台只读核对原件并生成报告。",
    outputs_desc="客户落责工作流对比报告.md → 来源搬运核对、三页月报逐格差异",
)
def monthly_luozhe_audit(
    stat_source: str = "", growth_source: str = "", process_book: str = "", report_ppt: str = "",
    batch: str = "", inbox: str | None = None, output_dir: str | None = None,
) -> TaskResult:
    root = Path(inbox) if inbox else INBOX_DIR
    target_dir = Path(output_dir) if output_dir else OUTPUT_DIR
    try:
        files = (_input(stat_source, ".xlsx", root), _input(growth_source, ".xlsx", root),
                 _input(process_book, ".xlsx", root), _input(report_ppt, ".pptx", root))
        if len(set(files)) != 4:
            raise ValueError("请为四个角色分别选择文件")
        source, growth, process = (load_workbook(f, read_only=False, data_only=True) for f in files[:3])
        deck = Presentation(files[3])
        if len(deck.slides) < 3:
            raise ValueError("月报至少需要三页客户落责表")
        for sheet in ("地市总体落责", "集团规上", "专精特新"):
            if sheet not in source:
                raise ValueError(f"落责来源表缺少工作表：{sheet}")
        for sheet in ("全量统计", "全量收入正增长统计", "集团规上统计", "专精特新统计"):
            if sheet not in process:
                raise ValueError(f"处理工作簿缺少工作表：{sheet}")

        source_steps = []
        for source_sheet, process_sheet, first_source, first_process, count in SOURCE_SHEETS:
            checked, differences = _compare_source(source, process, source_sheet, process_sheet,
                                                   first_source, first_process, count)
            source_steps.append((source_sheet, checked, differences))
        checked, differences = _compare_growth(growth, process)
        steps = [source_steps[0], ("收入正增长", checked, differences), *source_steps[1:]]
        material_source = [_compare_material_source(source, growth, process, config, i)
                           for i, config in enumerate(MATERIALS)]
        material = [_compare_material(process, deck, config, i) for i, config in enumerate(MATERIALS)]
        source.close(); growth.close(); process.close()
    except (ValueError, OSError, KeyError, IndexError, InvalidOperation) as exc:
        return TaskResult(ok=False, message=f"无法完成核对：{exc}")

    period = batch.strip() or "未确认"
    lines = ["# 客户落责月度工作流对比报告", "", f"账期标记：{period}（未从四份文件自动验证）", "",
             "四份输入文件均只读；本报告不修改处理表或月报。", "",
             "输入文件：" + "；".join(path.name for path in files), "",
             "## 7 步核对结果", "",
             "| 步骤 | 内容 | 核对结果 |", "| --- | --- | --- |"]
    names = ["更新全量统计", "更新全量收入正增长统计", "更新集团规上统计", "更新专精特新统计"]
    for i, ((label, checked, diff), step_name) in enumerate(zip(steps, names), 1):
        lines.append(f"| {i} | {step_name}：{label}来源 → 处理表数据源区 | {checked} 项，{len(diff)} 处差异 |")
    for i, (diff, labels) in enumerate(material, 5):
        lines.append(f"| {i} | 来源重算 → 处理表PPT材料 → 月报第{i - 4}页 | {23 * MATERIALS[i - 5][3]} 个数值；材料区 {len(material_source[i - 5])} 处差异，月报 {len(diff)} 处差异；{len(labels)} 处地市标签差异 |")
    lines += ["", "## 来源与处理表差异", ""]
    for label, _, diff in steps:
        if not diff:
            continue
        lines += [f"### {label}", "", "| 来源位置 | 处理表位置 | 来源值 | 处理表值 |", "| --- | --- | --- | --- |"]
        for item in diff:
            if len(item) == 6:
                ss, sc, ps, pc, a, b = item
                lines.append(f"| {ss}!{sc} | {ps}!{pc} | {a} | {b} |")
            else:
                sc, pc, a, b = item
                lines.append(f"| 正增长!{sc} | 全量收入正增长统计!{pc} | {a} | {b} |")
        lines.append("")
    if not any(diff for _, _, diff in steps):
        lines.append("四个来源搬运环节未发现差异。")
    lines += ["", "## 来源重算与处理表PPT材料差异", ""]
    for config, differences in zip(MATERIALS, material_source):
        lines.append(f"### {config[0]}（{len(differences)}处）")
        if differences:
            lines += ["", "| 地市 | 处理表单元格 | 来源重算值 | 处理表值 |",
                      "| --- | --- | ---: | ---: |"]
            for city, cell, value, actual in differences:
                lines.append(f"| {city} | {config[1]}!{cell} | {value} | {actual} |")
        else:
            lines.append("来源重算结果与处理表材料区一致。")
        lines.append("")
    lines += ["", "## 月报数值差异", ""]
    for page, ((diff, labels), config) in enumerate(zip(material, MATERIALS), 1):
        lines += [f"### 第{page}页：{config[0]}（{len(diff)}处）", "",
                  "| 地市 | 分级 | 指标 | 处理表单元格 | 处理表值 | 月报值 | 月报表格行 |",
                  "| --- | --- | --- | --- | ---: | ---: | ---: |"]
        for city, grade, metric, cell, value, ppt_value, slide_row in diff:
            lines.append(f"| {city} | {grade} | {metric} | {config[1]}!{cell} | {value} | {ppt_value} | {slide_row} |")
        if labels:
            lines += ["", "地市标签差异：" + "；".join(f"处理表{a}／月报{b}" for a, b, _, _ in labels)]
        lines.append("")
    lines += ["## 交付复核", "", "确认四份文件的实际统计期一致；核对三页表格后，按当期数值复核各页顶部总结。",
              "地市末行按月报模板显示：第1页“政企”，第2、3页“省管客户”；处理表内分别使用“省政企”和“省直管客户”。"]
    target_dir.mkdir(parents=True, exist_ok=True)
    out = target_dir / f"客户落责工作流对比报告-{datetime.now():%Y%m%d-%H%M%S}.md"
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    source_count = sum(len(diff) for _, _, diff in steps)
    material_count = sum(len(diff) for diff in material_source)
    ppt_count = sum(len(diff) for diff, _ in material)
    return TaskResult(ok=True, message=f"核对完成：来源搬运差异 {source_count} 处，材料区差异 {material_count} 处，月报数值差异 {ppt_count} 处",
                      outputs=[str(out)], metrics={"来源搬运差异": source_count, "材料区差异": material_count,
                                                  "月报数值差异": ppt_count},
                      warnings=["文件内统计期未自动验证，请确认四份文件属于同一批次"])
