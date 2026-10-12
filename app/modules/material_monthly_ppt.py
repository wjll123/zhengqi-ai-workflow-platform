"""大政企月度经营通报 P11–14 表格更新。仅替换原生表格数据。"""
from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook

from ..core.config import INBOX_DIR, OUTPUT_DIR, TEMPLATE_DIR, get_batch
from ..core.runner import TaskResult, registry

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
A = "{" + NS["a"] + "}"
SPECS = (
    (11, "双线三单  (月报用)", list(range(6)) + list(range(8, 20)), (5, 9, 13, 17)),
    (12, "双线到期 (月报用)", list(range(20)), (5, 7, 11, 13, 17, 19)),
    (13, "商企到期 (月报用)", list(range(11)), (4, 6, 8, 10)),
    (14, "团购到期 (月报)", list(range(11)), (4, 6, 8, 10)),
)


def _chosen(name: str, dirs: tuple[Path, ...], suffix: str) -> Path:
    if not name or name != Path(name).name or name.startswith(("~$", ".~")):
        raise ValueError("请从列表中选择文件")
    for directory in dirs:
        path = directory / name
        if path.is_file() and path.suffix.lower() == suffix:
            return path
    raise ValueError(f"文件不存在或格式不符：{name}")


def _cell_text(cell) -> str:
    return "".join(x.text or "" for x in cell.xpath(".//a:t", namespaces=NS)).strip()


def _put(cell, value: str, color: str | None = None) -> None:
    texts = cell.xpath(".//a:t", namespaces=NS)
    if not texts:
        paragraph = cell.find(".//" + A + "p")
        if paragraph is None:
            tx = cell.find(A + "txBody")
            if tx is None:
                tx = etree.SubElement(cell, A + "txBody")
            paragraph = etree.SubElement(tx, A + "p")
        run = etree.SubElement(paragraph, A + "r")
        texts = [etree.SubElement(run, A + "t")]
    texts[0].text = value
    for extra in texts[1:]:
        extra.text = ""
    if color:
        run = texts[0].getparent()
        if run.tag != A + "r":
            return
        props = run.find(A + "rPr")
        if props is None:
            props = etree.Element(A + "rPr")
            run.insert(0, props)
        fill = props.find(A + "solidFill")
        if fill is None:
            fill = etree.SubElement(props, A + "solidFill")
        for child in list(fill):
            fill.remove(child)
        etree.SubElement(fill, A + "srgbClr", val=color)


def _percent(value) -> str:
    return f"{(Decimal(str(value)) * 100).quantize(Decimal('0.1'), rounding=ROUND_HALF_UP):.1f}%"


def _rewrite_slide(xml: bytes, rows: list[tuple], cols: list[int], rates: tuple[int, ...], slide: int) -> bytes:
    root = etree.fromstring(xml)
    tables = root.xpath(".//a:tbl", namespaces=NS)
    if len(tables) != 1:
        raise ValueError(f"P{slide} 预期恰有一张原生表格")
    target_rows = tables[0].findall(A + "tr")
    if len(target_rows) != 24 or len(rows) < 25:
        raise ValueError(f"P{slide} 表格或源工作表行数不符")
    rankings = {}
    for target_col in rates:
        source_col = cols[target_col]
        vals = [(Decimal(str(rows[i + 1][source_col])), i) for i in range(3, 24)
                if isinstance(rows[i + 1][source_col], (int, float))]
        vals.sort(key=lambda item: item[0], reverse=True)
        rankings[target_col] = ({i for _, i in vals[:3]}, {i for _, i in vals[-3:]})
    for target_row in range(2, 24):
        source = rows[target_row + 1]
        cells = target_rows[target_row].findall(A + "tc")
        if len(cells) != len(cols):
            raise ValueError(f"P{slide} 第 {target_row + 1} 行列数不符")
        if _cell_text(cells[1]) != str(source[1]).strip():
            raise ValueError(f"P{slide} 地市顺序不符：PPT={_cell_text(cells[1])}，源={source[1]}")
        for target_col in range(2, len(cols)):
            value = source[cols[target_col]]
            if value is None or value == "":
                raise ValueError(f"P{slide} {source[1]} 第 {target_col + 1} 列缺值")
            if target_col in rates and isinstance(value, (int, float)):
                display = _percent(value)
            elif isinstance(value, (int, float)) and float(value).is_integer():
                display = str(int(value))
            else:
                display = str(value)
            color = "000000" if target_row >= 3 and target_col in rates else None
            if target_row >= 3 and target_col in rankings:
                blue, red = rankings[target_col]
                color = "4472C4" if target_row in blue else "C00000" if target_row in red else "000000"
            _put(cells[target_col], display, color)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


@registry.register(
    key="monthly_report_ppt_tables",
    name="大政企月度经营通报表格更新（P11–14）",
    module="材料输出",
    description="选择已有或新上传的周报数据和 PPT 模板，更新 P11–14 的表格",
    source_type="auto",
    source_note="选择周报工作簿的四张“月报用”工作表，并选择同账期 PPT 模板；平台只更新 P11–14 的表格数据。",
    outputs_desc="大政企月度经营通报表格更新-账期.pptx → P11–14 表格",
    params=[{"name": "source_name", "label": "周报数据源文件", "type": "file", "required": True},
            {"name": "template_name", "label": "PPT 模板文件", "type": "file", "required": True}],
)
def monthly_report_ppt_tables(source_name: str = "", template_name: str = "") -> TaskResult:
    try:
        source = _chosen(source_name, (INBOX_DIR,), ".xlsx")
        template = _chosen(template_name, (TEMPLATE_DIR, INBOX_DIR), ".pptx")
        batch = get_batch()
        book = load_workbook(source, data_only=True, read_only=True)
        first = book[SPECS[0][1]].cell(1, 1).value
        matched = re.search(r"(20\d{4})\d{2}", str(first))
        if not matched or matched.group(1) != batch:
            raise ValueError(f"源工作簿账期与当前账期 {batch} 不一致")
        changed = {}
        with ZipFile(template) as package:
            cover = etree.fromstring(package.read("ppt/slides/slide1.xml"))
            cover_text = "".join(x.text or "" for x in cover.xpath(".//a:t", namespaces=NS))
            year, month = int(batch[:4]), int(batch[4:6])
            if f"{year}年{month}月" not in cover_text:
                raise ValueError(f"PPT 模板封面月份与当前账期 {batch} 不一致")
            for slide, sheet, cols, rates in SPECS:
                if sheet not in book:
                    raise ValueError(f"源工作簿缺少工作表：{sheet}")
                rows = list(book[sheet].values)
                path = f"ppt/slides/slide{slide}.xml"
                changed[path] = _rewrite_slide(package.read(path), rows, cols, rates, slide)
            output = OUTPUT_DIR / f"大政企月度经营通报表格更新-{batch}.pptx"
            stream = BytesIO()
            with ZipFile(stream, "w") as result:
                for item in package.infolist():
                    result.writestr(item, changed.get(item.filename, package.read(item.filename)))
        book.close()
        output.write_bytes(stream.getvalue())
    except (ValueError, KeyError, OSError) as exc:
        return TaskResult(ok=False, message=str(exc))
    return TaskResult(ok=True, message="P11–14 共 4 张表格更新完成", outputs=[str(output)],
                      metrics={"数据源": source.name, "模板": template.name, "账期": batch},
                      warnings=["P3–10、点评文字与非表格页面沿用所选模板，请按对应账期核对。"])
