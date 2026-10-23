import json
import threading
import time
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from flask import Flask
from openpyxl import Workbook, load_workbook

from app.qcc.errors import BrowserActionRequired, CompanyNotFound, QueryBusy, QueryError, QueryPaused
from app.qcc.extraction import company_data, search_target
from app.qcc.jobs import QueryJobs
from app.qcc.safari import SafariSession
from app.qcc.sources import read_codes, read_names
from app.qcc.workbooks import export_job, read_previous
from app.qcc.pacing import NavigationPacer
from app.qcc.fields import SCHEMA_VERSION

URL = 'https://m.qcc.com/firm/' + 'a' * 32 + '.html'


def page(company):
    return 'window.__INITIAL_STATE__=' + json.dumps({'company': {'companyDetail': company}}, ensure_ascii=False)


def company(name='甲有限公司'):
    return {'Name': name, 'KeyNo': 'a' * 32, 'CreditCode': 'A12345678901234567',
            'Partners': [{'StockName': '甲股东', 'StockPercentValue': 60}],
            'Employees': [{'Name': '甲人员', 'Job': '董事'}],
            'CountInfo': {'PartnersCount': 1, 'EmployeesCount': 1}}


def data(name='甲有限公司'):
    return company_data(page(company(name)), name, url=URL)


class ParsingTests(unittest.TestCase):
    def test_all_fields_installments_zero_and_source_evidence_are_preserved(self):
        value = company()
        value['Partners'][0].update(ShouldCapi=0, RealCapi=0, InvestType='货币', CapiDate='2025-12-31',
            Tags=['大股东'], FinalBenefitPercent='60%', NewDisclosure='新增业务资料',
            PartnerShouldDetailList=[{'ShouldCapi':0,'ShoudDate':'2026-01-01','InvestType':'货币'},
                                    {'ShouldCapi':'100','ShoudDate':'2027-01-01','InvestType':'货币'}],
            PartnerRealDetailList=[{'RealCapi':0,'CapiDate':'2025-12-31','InvestName':None}])
        value['Employees'][0].update(CompanyCount=0, Tags=['实际控制人'], HasIntro=1, CerNo='123456789012345678')
        result = company_data(page(value), '甲有限公司', url=URL)
        mapped = dict(zip(result['shareholders']['headers'], result['shareholders']['rows'][0]))
        self.assertEqual(mapped['认缴出资（万元）'], 0)
        self.assertEqual(mapped['其他披露字段·NewDisclosure'], '新增业务资料')
        self.assertEqual(mapped['实缴出资日期'], '2025-12-31')
        self.assertEqual(len(result['contributions']['rows']), 3)
        self.assertEqual(result['contributions']['rows'][2][3], 0)
        self.assertEqual(result['contributions']['rows'][2][5], '未披露')
        self.assertEqual(result['raw_records']['people'][0]['CerNo'], '123********5678')
        self.assertIn('人物简介有入口', result['additional_details'])
        self.assertEqual(result['coverage']['shareholders'], {'expected':1,'captured':1})

    def test_malformed_installments_cannot_be_marked_complete(self):
        value = company(); value['Partners'][0]['PartnerRealDetailList'] = ['truncated']
        with self.assertRaisesRegex(QueryError, '分期出资'):
            company_data(page(value), '甲有限公司', url=URL)

    def test_identity_counts_and_numeric_percentage(self):
        result = data()
        self.assertEqual(result['shareholders']['rows'][0][2], 0.6)
        self.assertEqual(result['credit_code'], 'A12345678901234567')

    def test_incomplete_details_are_rejected(self):
        value = company()
        value['CountInfo']['PartnersCount'] = 2
        with self.assertRaisesRegex(QueryError, '应有 2 条'):
            company_data(page(value), '甲有限公司', url=URL)

    def test_zero_records_are_distinct_from_missing(self):
        value = company()
        value['Partners'] = None
        value['CountInfo']['PartnersCount'] = '0'
        self.assertEqual(company_data(page(value), '甲有限公司', url=URL)['shareholders']['rows'], [])
        value['CountInfo'].pop('PartnersCount')
        with self.assertRaises(QueryError):
            company_data(page(value), '甲有限公司', url=URL)

    def test_wrong_identity_is_rejected(self):
        for name, code in [('乙有限公司', ''), ('甲有限公司', 'B12345678901234567')]:
            with self.subTest(name=name, code=code), self.assertRaises(QueryError):
                company_data(page(company()), name, code, URL)

    def test_verified_former_name_and_institution_key(self):
        value = company('新名称研究院')
        value['KeyNo'] = 'g' + 'a' * 32
        value['OriginalName'] = [{'Name': '旧名称研究院'}]
        value['Partners'] = []
        value['Employees'] = []
        value['CountInfo'] = {'PartnersCount': 0, 'EmployeesCount': 0}
        url = 'https://www.qcc.com/firm/g' + 'a' * 32 + '.html'
        result = company_data(page(value), '旧名称研究院', url=url)
        self.assertEqual(result['coverage']['people'], {'expected': 0, 'captured': 0})
        self.assertIn('原名称与名单一致', result['match_note'])
        with self.assertRaisesRegex(QueryError, '名称与名单不一致'):
            company_data(page(value), '其他研究院', url=url)

    def test_verification_on_redirect_and_search_is_detected(self):
        for source, url in [('<title>用户验证</title>', URL), ('', 'https://m.qcc.com/mlogin'), ('', 'https://www.qcckyc.com/')]:
            with self.subTest(url=url), self.assertRaises(BrowserActionRequired):
                company_data(source, '甲有限公司', url=url)

    def test_search_uses_credit_code_for_duplicate_names(self):
        results = [{'Name':'甲有限公司', 'KeyNo':'a'*32, 'CreditCode':'A12345678901234567'},
                   {'Name':'甲有限公司', 'KeyNo':'b'*32, 'CreditCode':'B12345678901234567'}]
        source = 'window.__INITIAL_STATE__=' + json.dumps({'search':{'searchRes':{'Status':200,'Result':results}}})
        with self.assertRaises(QueryError):
            search_target(source, '甲有限公司')
        self.assertEqual(search_target(source, '甲有限公司', 'A12345678901234567'), URL)

    def test_missing_ratio_and_invalid_numbers_fail(self):
        for ratio in (None, -1, 101, float('nan'), True):
            value = company()
            value['Partners'][0]['StockPercentValue'] = ratio
            with self.subTest(ratio=ratio), self.assertRaises(QueryError):
                company_data(page(value), '甲有限公司', url=URL)

    def test_credit_code_fallback_checks_search_and_detail_and_records_name(self):
        code = 'A12345678901234567'
        source = 'window.__INITIAL_STATE__=' + json.dumps({'search': {'searchRes': {'Status': 200, 'Result': [
            {'Name': '新名称有限公司', 'KeyNo': 'a'*32, 'CreditCode': code}]}}})
        with self.assertRaises(CompanyNotFound):
            search_target(source, '旧名称有限公司', code)
        self.assertEqual(search_target(source, '旧名称有限公司', code, by_code=True), URL)
        result = company_data(page(company('新名称有限公司')), '旧名称有限公司', code, URL, by_code=True)
        self.assertEqual(result['company'], '新名称有限公司')
        self.assertIn('原始名单名称已保留', result['match_note'])
        for bad_code in ('', 'B12345678901234567'):
            with self.subTest(code=bad_code), self.assertRaises(QueryError):
                company_data(page(company('新名称有限公司')), '旧名称有限公司', bad_code, URL, by_code=True)

    def test_code_search_still_rejects_conflicting_results(self):
        source = 'window.__INITIAL_STATE__=' + json.dumps({'search': {'searchRes': {'Status': 200, 'Result': [
            {'Name': '名称有限公司', 'KeyNo': 'a'*32, 'CreditCode': 'B12345678901234567'}]}}})
        with self.assertRaises(CompanyNotFound):
            search_target(source, '名称有限公司', 'A12345678901234567', by_code=True)


class SourceTests(unittest.TestCase):
    def test_read_deduplicate_and_limit(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            (root/'names.csv').write_text('公司名称\n甲（广东）有限公司\n甲(广东)有限公司\n乙有限公司\n')
            self.assertEqual(len(read_names('names.csv', root)), 2)
            (root/'names.csv').write_text('公司名称\n'+'\n'.join(f'企业{i}有限公司' for i in range(101)))
            with self.assertRaisesRegex(QueryError, '101'):
                read_names('names.csv', root)
            with self.assertRaises(QueryError):
                read_names('../names.csv', root)

    def test_incremental_and_identity_conflict(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            book = Workbook(); sheet = book.active
            sheet.append([None,'企业名称']); sheet.append(['旧有限公司','旧有限公司']); sheet.append([None,'新有限公司'])
            book.save(root/'names.xlsx'); book.close()
            self.assertEqual(read_names('names.xlsx', root), ['新有限公司'])
            (root/'codes.csv').write_text('公司名称,统一社会信用代码\n新有限公司,A12345678901234567\n新有限公司,B12345678901234567\n')
            with self.assertRaisesRegex(QueryError, '多个信用代码'):
                read_codes('codes.csv', ['新有限公司'], root)


class JobTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(); self.root = Path(self.temp.name)
        self.inbox = self.root/'inbox'; self.inbox.mkdir()
        (self.inbox/'names.csv').write_text('公司名称\n甲有限公司\n乙有限公司\n丙有限公司\n')
        self.calls = []; self.block = True; self.sessions = []
        outer = self
        class Session:
            def __init__(self, window_id=None):
                self.window_id = window_id or '123'; self.preserve_window = False; outer.sessions.append(self)
            def __enter__(self): return self
            def __exit__(self, *_args): pass
            def query(self, name, credit_code='', detail_url=''):
                outer.calls.append(name)
                if name == '乙有限公司' and outer.block: raise BrowserActionRequired('请完成验证')
                return data(name)
        self.jobs = QueryJobs(self.root/'state', self.inbox, self.root/'outputs', Session)
        self.wait_patch = patch.object(self.jobs, '_wait', return_value=None); self.wait_patch.start()

    def tearDown(self):
        self.wait_patch.stop(); self.temp.cleanup()

    def finish(self, job_id):
        end = time.monotonic()+5
        while time.monotonic()<end:
            job = self.jobs.get(job_id)
            if job['status'] != 'running': return job
            time.sleep(.01)
        self.fail('worker did not finish')

    def start(self, **kwargs):
        return self.jobs.start({'source_name':'names.csv', **kwargs})

    def test_block_saves_progress_preserves_window_and_resumes_only_remaining(self):
        job = self.finish(self.start()['id'])
        self.assertEqual(job['status'], 'blocked'); self.assertEqual(job['counts']['success'], 1)
        self.assertTrue(self.sessions[0].preserve_window)
        snapshot = QueryJobs(self.root/'state', self.inbox, self.root/'outputs').get(job['id'])
        self.assertEqual(snapshot['counts']['success'], 1)
        self.block = False
        self.jobs.resume(job['id']); done = self.finish(job['id'])
        self.assertEqual(done['counts']['success'], 3)
        self.assertEqual(self.calls, ['甲有限公司','乙有限公司','乙有限公司','丙有限公司'])

    def test_limit_and_concurrent_start(self):
        job = self.finish(self.start(max_new_queries=1)['id'])
        self.assertEqual(job['status'], 'paused'); self.assertEqual(self.calls, ['甲有限公司'])
        lock = self.jobs._lock()
        try:
            with self.assertRaises(QueryBusy): self.start()
        finally: lock.close()

    def test_legacy_success_is_refreshed_and_preserved_until_replaced(self):
        job = self.finish(self.start(max_new_queries=1)['id'])
        raw = self.jobs._read(job['id']); raw['rows'][0]['data'].pop('schema_version')
        self.jobs._save(raw)
        self.jobs.resume(job['id']); done = self.finish(job['id'])
        self.assertEqual(self.calls, ['甲有限公司', '甲有限公司'])
        self.assertEqual(done['rows'][0]['data']['schema_version'], SCHEMA_VERSION)
        self.assertIn('previous_data', self.jobs._read(job['id'])['rows'][0])
        self.assertNotIn('previous_data', done['rows'][0])
        self.assertNotIn('raw_records', done['rows'][0]['data'])

    def test_retained_old_rows_remain_in_export_without_claiming_completion(self):
        job = self.finish(self.start(max_new_queries=1)['id'])
        raw = self.jobs._read(job['id']); raw['rows'][0]['data'].pop('schema_version')
        self.jobs._prepare_full(raw); self.jobs._save(raw)
        path = self.jobs.export(job['id'])
        book = load_workbook(path, data_only=True)
        self.assertEqual(book['查询状态']['B2'].value, '待查询')
        self.assertIn('旧版基础字段', book['查询状态']['J2'].value)
        self.assertEqual(book['股东信息'].max_row, 2)
        book.close()
        restored = read_previous(path.name, ['甲有限公司','乙有限公司','丙有限公司'],path.parent)
        self.assertEqual(restored['甲有限公司']['status'], 'pending')

    def test_recover_interrupted_query(self):
        job = self.finish(self.start(max_new_queries=1)['id'])
        raw = self.jobs._read(job['id']); raw['status']='running'; raw['rows'][1]['status']='querying'
        self.jobs._save(raw)
        restored = self.jobs.get(job['id'])
        self.assertEqual(restored['status'], 'paused')
        self.assertEqual(restored['rows'][1]['status'], 'pending')
        self.assertEqual(restored['counts']['success'], 1)

    def test_pause_is_observed_during_wait(self):
        self.block = False
        with patch.object(self.jobs, '_wait', side_effect=lambda job:self.jobs.pause(job['id'])):
            job = self.finish(self.start()['id'])
        self.assertEqual(job['status'], 'paused'); self.assertEqual(self.calls, ['甲有限公司'])

    def test_export_round_trip_preserves_unknown_and_numeric_values(self):
        job = self.finish(self.start()['id']); path = self.jobs.export(job['id'])
        book = load_workbook(path, data_only=True)
        self.assertEqual(book['股东信息']['D2'].value, .6)
        self.assertEqual(book['查询状态']['B3'].value, '待验证')
        self.assertEqual(book['查询状态']['H2'].number_format, '@')
        self.assertIsNone(book['查询状态']['C4'].value)
        book.close()
        restored = read_previous(path.name, ['甲有限公司','乙有限公司','丙有限公司'],path.parent)
        self.assertEqual(list(restored), ['甲有限公司'])
        self.assertEqual(restored['甲有限公司']['data']['schema_version'], SCHEMA_VERSION)
        self.assertIn('contributions', restored['甲有限公司']['data'])
        with self.assertRaises(QueryError): read_previous(path.name, ['其他有限公司'], path.parent)

    def test_invalid_settings_do_not_start_browser(self):
        for params in ({'max_new_queries':True},{'max_new_queries':101},{'delay_seconds':0},{'delay_seconds':float('nan')}, {'source_name':[]}):
            with self.subTest(params=params), self.assertRaises(QueryError): self.start(**params)
        self.assertEqual(self.calls, [])

    def test_api_validation_and_async_start(self):
        from app.qcc import routes
        app=Flask(__name__); app.register_blueprint(routes.bp)
        with patch.object(routes,'jobs',self.jobs), app.test_client() as client:
            self.assertEqual(client.post('/api/qcc/jobs',json=[]).status_code,400)
            self.assertEqual(client.get('/api/qcc/jobs/unknown').status_code,400)
            response=client.post('/api/qcc/jobs',json={'source_name':'names.csv','max_new_queries':1})
            self.assertEqual(response.status_code,202)
            job=self.finish(response.json['job']['id'])
            with client.get('/api/qcc/jobs/'+job['id']+'/export') as download:
                self.assertEqual(download.status_code,200)


class SafariTests(unittest.TestCase):
    def test_credit_code_uses_only_one_search(self):
        session = SafariSession('123')
        with patch.object(session, '_page', return_value=('about:blank', '')), patch.object(session, '_load', side_effect=[
            URL, data()]) as load:
            result = session.query('甲有限公司', 'A12345678901234567')
        self.assertEqual(result['company'], '甲有限公司')
        self.assertEqual(load.call_count, 2)
        self.assertEqual(load.call_args_list[0].args[0], 'https://www.qcc.com/web/search?key=A12345678901234567')

    def test_known_url_skips_search_and_checks_identity(self):
        session = SafariSession('123')
        with patch.object(session, '_page', return_value=('about:blank', '')), patch.object(session, '_load', return_value=data()) as load:
            session.query('甲有限公司', 'A12345678901234567', detail_url=URL)
            self.assertEqual(load.call_count, 1)
            self.assertEqual(load.call_args.args[0], URL)
            wrong = company('错误企业'); wrong['CreditCode'] = 'B12345678901234567'
            with self.assertRaises(QueryError):
                load.call_args.args[1](page(wrong), URL)

    def test_stop_during_rate_wait_prevents_navigation(self):
        session = SafariSession('123')
        session.before_navigation = lambda: (_ for _ in ()).throw(QueryPaused('停止'))
        with patch('app.qcc.safari.apple') as apple, self.assertRaises(QueryPaused):
            session._load(URL, lambda *args: None)
        apple.assert_not_called()

    def test_blocked_window_is_not_closed(self):
        session=SafariSession('123')
        with patch('app.qcc.safari.apple') as apple:
            session.__exit__(BrowserActionRequired, BrowserActionRequired('验证'), None)
            apple.assert_not_called()

    def test_previous_search_query_is_not_used(self):
        session=SafariSession('123')
        source='window.__INITIAL_STATE__={}'
        with patch('app.qcc.safari.apple'), patch('app.qcc.safari.time.sleep'), patch.object(session,'_page',side_effect=[
            ('https://www.qcc.com/web/search?key=old',source),('https://www.qcc.com/web/search?key=new',source)]):
            result=session._load('https://www.qcc.com/web/search?key=new',lambda s,u:u)
        self.assertTrue(result.endswith('key=new'))

    def test_verified_detail_is_reused(self):
        session=SafariSession('123')
        with patch.object(session,'_page',return_value=(URL,page(company()))),patch.object(session,'_load') as load:
            self.assertEqual(session.query('甲有限公司')['company'],'甲有限公司'); load.assert_not_called()


class PacingTests(unittest.TestCase):
    def test_every_navigation_is_spaced_and_tenth_page_gets_rest(self):
        with TemporaryDirectory() as temp:
            clock = [1000.0]
            def sleep(seconds): clock[0] += seconds
            path = Path(temp)/'rate.json'
            pacer = NavigationPacer(path, clock=lambda:clock[0], sleep=sleep)
            times = []
            for _ in range(11):
                pacer.before_navigation(); times.append(clock[0])
            self.assertEqual(times[1] - times[0], 60)
            self.assertEqual(times[10] - times[9], 300)
            restored = NavigationPacer(path, clock=lambda:clock[0], sleep=sleep)
            restored.before_navigation()
            self.assertEqual(clock[0] - times[-1], 60)

    def test_verification_cooldown_and_slower_interval_survive_restart(self):
        with TemporaryDirectory() as temp:
            clock = [1000.0]
            def sleep(seconds): clock[0] += seconds
            path = Path(temp)/'rate.json'
            pacer = NavigationPacer(path, clock=lambda:clock[0], sleep=sleep)
            pacer.before_navigation(); pacer.blocked()
            restored = NavigationPacer(path, clock=lambda:clock[0], sleep=sleep)
            restored.before_navigation(); self.assertEqual(clock[0], 1900)
            restored.before_navigation(); self.assertEqual(clock[0], 2020)
            restored.blocked()
            self.assertEqual(restored.state['penalty_interval'], 240)
            self.assertEqual(restored.state['cooldown_until'], 3820)


if __name__ == '__main__': unittest.main()
