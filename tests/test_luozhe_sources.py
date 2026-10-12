import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app.core.config import TEMPLATE_DIR
from app.core.runner import Task, TaskResult, registry
from app.modules import load_all
from app.modules.data_luozhe import (
    CITY_ORDER, GRADES, STAT_SHEETS, TEMPLATE_FILE, SourceDataError, _fill_template_workbook, _find,
    _validate_output_template,
    _validate_grow_sheet, _validate_stat_sheet, monthly_luozhe,
)


class LuozheSourceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        load_all()

    def test_uploaded_files_without_period_are_ready_and_latest_is_used(self):
        with tempfile.TemporaryDirectory() as tmp:
            inbox = Path(tmp)
            stat = "【数据来源-月度更新】21各地市落责情况统计及排名（地市总体、集团规上、专精特新）"
            growth = "【数据来源-月度更新】 月度收入正增长客户数量"
            names = [f"{stat}.xlsx", f"{stat}(1).xlsx", f"{growth}.xlsx", f"{growth}(1).xlsx"]
            for index, name in enumerate(names):
                path = inbox / name
                path.touch()
                os.utime(path, (100 + index, 100 + index))

            status = registry.get("monthly_luozhe").check_sources(inbox, batch="202609")
            self.assertTrue(status["ready"])
            self.assertEqual([item["file"] for item in status["items"]], [names[1], names[3]])
            self.assertEqual(_find(inbox, "21各地市落责", "落责情况统计").name, names[1])
            self.assertEqual(_find(inbox, "收入正增长", "正增长客户数量").name, names[3])

    def test_non_workbook_does_not_satisfy_luozhe_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            inbox = Path(tmp)
            (inbox / "21各地市落责情况统计.txt").touch()
            (inbox / "月度收入正增长客户数量.xlsx").touch()
            status = registry.get("monthly_luozhe").check_sources(inbox, batch="202609")
            self.assertFalse(status["ready"])
            self.assertFalse(status["items"][0]["found"])
            self.assertTrue(status["items"][1]["found"])

    def test_other_tasks_can_still_require_period_in_filename(self):
        task = Task(key="period_test", name="测试", module="测试", description="",
                    handler=lambda: TaskResult(ok=True),
                    required_sources=[{"label": "当期来源", "keywords": ["来源"], "batch_key": True}])
        with tempfile.TemporaryDirectory() as tmp:
            inbox = Path(tmp)
            (inbox / "来源.xlsx").touch()
            self.assertFalse(task.check_sources(inbox, batch="202609")["ready"])
            (inbox / "来源-202609.xlsx").touch()
            self.assertTrue(task.check_sources(inbox, batch="202609")["ready"])


def stat_rows(grades):
    heading_city = [None] * 19
    heading_city[0] = "地市"
    heading_city[17:19] = ["落责不准确", "未落责"]
    heading_grade = [None] * 19
    heading_grade[1] = "客户分级"
    heading_grade[2] = "落责准确率"
    heading_grade[3], heading_grade[6], heading_grade[9] = "行客", "商客", "综合营服"
    heading_grade[12] = "总计"
    heading_values = [None] * 19
    for col, label in {3: "落责数量", 4: "直销落责数量", 5: "有客户经理数",
                       6: "落责数量", 7: "直销落责数量", 8: "有客户经理数",
                       9: "落责数量", 10: "直销落责数量", 11: "有客户经理数",
                       12: "落责数量", 13: "直销落责数量", 14: "其中：有客户经理数",
                       15: "落责准确数", 16: "有客户经理占比"}.items():
        heading_values[col] = label
    rows = [["说明"], heading_city, heading_grade, heading_values]
    for city in ["全省（维度一：产权与注册地判断）",
                 "全省（维度二：产权对应地市是否有实际业务收入判断）", *CITY_ORDER]:
        for index, grade in enumerate(grades):
            row = [0] * 19
            row[0] = city if index == 0 else None
            row[1], row[2] = grade, 0.5
            row[12:17] = [10, 10, 5, 5, 0.5]
            rows.append(row)
    return rows


def grow_rows():
    return [["序号", "客户所属地市", *GRADES]] + [
        [index, city, 0, 1, 2, 3, 4]
        for index, city in enumerate(["全省", *CITY_ORDER])
    ]


class LuozheValidationTests(unittest.TestCase):
    def test_zero_is_valid_but_missing_number_fails(self):
        rows = grow_rows()
        self.assertEqual(_validate_grow_sheet(rows)["深圳"]["卓越"], 0)
        rows[2][2] = None
        with self.assertRaisesRegex(SourceDataError, "深圳/卓越.*缺少有效数值"):
            _validate_grow_sheet(rows)
        rows[2][2] = "/"
        with self.assertRaisesRegex(SourceDataError, "深圳/卓越.*缺少有效数值"):
            _validate_grow_sheet(rows)

    def test_wrong_header_and_missing_city_grade_fail(self):
        rows = stat_rows(GRADES)
        self.assertIn("深圳", _validate_stat_sheet(rows, "地市总体落责"))
        rows[3][13] = "错误列"
        with self.assertRaisesRegex(SourceDataError, "模板不匹配.*直销落责数量"):
            _validate_stat_sheet(rows, "地市总体落责")
        rows[3][13] = "直销落责数量"
        rows[4 + 2 * len(GRADES)][1] = "错误分级"
        with self.assertRaisesRegex(SourceDataError, "深圳/卓越"):
            _validate_stat_sheet(rows, "地市总体落责")

    def test_missing_stat_value_and_wrong_growth_template_fail(self):
        stat = stat_rows(GRADES)
        stat[4 + 2 * len(GRADES)][12] = None
        with self.assertRaisesRegex(SourceDataError, "深圳/卓越/落责数量.*缺少有效数值"):
            _validate_stat_sheet(stat, "地市总体落责")
        stat[4 + 2 * len(GRADES)][12] = "/"
        with self.assertRaisesRegex(SourceDataError, "深圳/卓越/落责数量.*缺少有效数值"):
            _validate_stat_sheet(stat, "地市总体落责")
        stat[4 + 2 * len(GRADES)][12] = 0
        stat[4 + 2 * len(GRADES)][2] = "/"
        self.assertIn("深圳", _validate_stat_sheet(stat, "地市总体落责"))
        stat[4 + 2 * len(GRADES)][12] = 10
        with self.assertRaisesRegex(SourceDataError, "深圳/卓越/落责准确率.*缺少有效数值"):
            _validate_stat_sheet(stat, "地市总体落责")

        growth = grow_rows()
        growth[0][1] = "其他字段"
        with self.assertRaisesRegex(SourceDataError, "收入正增长文件模板不匹配"):
            _validate_grow_sheet(growth)

    def test_missing_province_dimension_and_raw_field_fail(self):
        rows = stat_rows(GRADES)
        rows[4][0] = "其他维度"
        with self.assertRaisesRegex(SourceDataError, "全省维度一"):
            _validate_stat_sheet(rows, "地市总体落责")
        rows = stat_rows(GRADES)
        rows[4][3] = None
        with self.assertRaisesRegex(SourceDataError, "第 4 列.*缺少有效数值"):
            _validate_stat_sheet(rows, "地市总体落责")

    def test_missing_required_sheet_stops_before_writing_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            inbox = Path(tmp)
            (inbox / "21各地市落责情况统计.xlsx").touch()
            (inbox / "月度收入正增长客户数量.xlsx").touch()
            with patch("app.modules.data_luozhe.tables.sheet_names", return_value=["地市总体落责"]), \
                 patch("app.modules.data_luozhe.load_workbook") as write:
                result = monthly_luozhe(str(inbox))
            self.assertFalse(result.ok)
            self.assertIn("集团规上", result.message)
            self.assertIn("专精特新", result.message)
            write.assert_not_called()

    def test_bad_source_stops_before_writing_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            inbox = Path(tmp)
            (inbox / "21各地市落责情况统计.xlsx").touch()
            (inbox / "月度收入正增长客户数量.xlsx").touch()
            stat = {name: stat_rows(grades) for name, grades in STAT_SHEETS.items()}
            growth = grow_rows()
            growth[2][2] = None

            def read_rows(_path, sheet):
                return stat[sheet] if sheet else growth

            with patch("app.modules.data_luozhe.tables.sheet_names", return_value=list(STAT_SHEETS)), \
                 patch("app.modules.data_luozhe._rows_of", side_effect=read_rows), \
                 patch("app.modules.data_luozhe.load_workbook") as write:
                result = monthly_luozhe(str(inbox))
            self.assertFalse(result.ok)
            self.assertIn("深圳/卓越", result.message)
            write.assert_not_called()


class LuozheTemplateTests(unittest.TestCase):
    def test_missing_template_region_fails(self):
        wb = load_workbook(TEMPLATE_DIR / TEMPLATE_FILE)
        wb["全量收入正增长统计"]["AV4"] = None
        with self.assertRaisesRegex(SourceDataError, "AV4.*PPT材料"):
            _validate_output_template(wb)
        wb.close()

    def test_output_keeps_template_layout_and_fills_all_regions(self):
        template = load_workbook(TEMPLATE_DIR / TEMPLATE_FILE)
        _validate_output_template(template)
        raw_stat = {sheet: stat_rows(grades) for sheet, grades in STAT_SHEETS.items()}
        parsed = {sheet: _validate_stat_sheet(rows, sheet) for sheet, rows in raw_stat.items()}
        raw_grow = grow_rows()
        grow = _validate_grow_sheet(raw_grow)
        output = load_workbook(TEMPLATE_DIR / TEMPLATE_FILE)
        _fill_template_workbook(output, raw_stat, raw_grow, grow,
                                parsed["地市总体落责"], parsed["集团规上"], parsed["专精特新"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "output.xlsx"
            output.save(path)
            output.close()
            result = load_workbook(path)
            self.assertEqual(result.sheetnames, template.sheetnames)
            for expected, actual in zip(template, result):
                self.assertEqual((actual.max_row, actual.max_column), (expected.max_row, expected.max_column))
                self.assertEqual(set(map(str, actual.merged_cells.ranges)), set(map(str, expected.merged_cells.ranges)))
                self.assertEqual(len(actual.conditional_formatting), len(expected.conditional_formatting))
                self.assertFalse(any(cell.data_type == "f" for row in actual for cell in row))
                self.assertTrue(all(expected[cell.coordinate]._style == cell._style
                                    for row in actual for cell in row))
            self.assertEqual(result["全量统计"]["C5"].value, 10)
            self.assertEqual(result["全量统计"]["O33"].value, 10)
            self.assertEqual(result["全量收入正增长统计"]["C34"].value, 10)
            self.assertEqual(result["全量收入正增长统计"]["AL34"].value, 0)
            self.assertEqual(result["全量收入正增长统计"]["AW7"].value, 10)
            self.assertEqual(result["集团规上统计"]["C6"].value, 10)
            self.assertEqual(result["集团规上统计"]["AB7"].value, 10)
            self.assertEqual(result["专精特新统计"]["O35"].value, 10)
            result.close()
        template.close()


if __name__ == "__main__":
    unittest.main()
